from pathlib import Path

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from make_samples import make_all, _band, HEADER_BOTTOM, FOOTER_TOP, SLIDE_W
from placement import Box, Entry, build_deck, default_layout, fit_top_left, image_area, obstructions, FOOTER_TYPES, TITLE_TYPES
from potx import open_template


@pytest.fixture(scope="module")
def samples():
    return make_all()


def _layout(template, name):
    prs = open_template(template)
    return prs, next(l for l in prs.slide_layouts if l.name == name)


def test_default_image_area_is_full_width_between_header_and_footers(samples):
    template, _ = samples
    prs, layout = _layout(template, "PBI Page")
    assert image_area(layout, prs.slide_width, prs.slide_height) == Box(0, HEADER_BOTTOM, SLIDE_W, FOOTER_TOP - HEADER_BOTTOM)


def test_picture_placeholder_overrides_image_area(samples):
    template, _ = samples
    prs, layout = _layout(template, "PBI Page (Picture)")
    pic = next(ph for ph in layout.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
    assert image_area(layout, prs.slide_width, prs.slide_height) == Box.of(pic)


def test_fit_is_height_limited_and_anchored_top_left():
    area = Box(0, 1000, 12_000, 5_000)
    box = fit_top_left(area, 1600, 900)
    assert (box.left, box.top, box.height) == (0, 1000, 5_000)
    assert box.width < area.width


def test_fit_uses_full_width_when_height_allows():
    area = Box(0, 1000, 12_000, 5_000)
    box = fit_top_left(area, 3200, 900)
    assert (box.left, box.top, box.width) == (0, 1000, 12_000)
    assert box.height <= area.height


def test_fit_never_distorts():
    box = fit_top_left(Box(0, 0, 12_000, 5_000), 1920, 1440)
    assert box.width / box.height == pytest.approx(1920 / 1440, rel=1e-3)


@pytest.mark.parametrize("layout_name", ["PBI Page", "PBI Page (Picture)"])
def test_deck_keeps_template_slides_and_never_covers_header_or_footers(samples, tmp_path, layout_name):
    template, pages = samples
    entries = [Entry(title=p.stem, image=p) for p in pages] + [Entry(title="broken", image=None, error="x")]
    out = build_deck(template, entries, tmp_path / "deck.pptx", layout_name=layout_name)

    prs = Presentation(out)
    slides = list(prs.slides)
    assert len(slides) == len(entries)  # the template's cover slide is dropped

    for slide, entry in zip(slides, entries):
        assert slide.shapes.title.text == entry.title
        kinds = {ph.placeholder_format.type for ph in slide.placeholders}
        assert kinds <= TITLE_TYPES | FOOTER_TYPES
        footers = [Box.of(ph) for ph in slide.placeholders if ph.placeholder_format.type in FOOTER_TYPES]
        assert len(footers) == 3
        header = Box(0, 0, SLIDE_W, HEADER_BOTTOM)
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if entry.image is None:
            assert not pictures
            continue
        (pic,) = pictures
        box = Box.of(pic)
        area = image_area(slide.slide_layout, prs.slide_width, prs.slide_height)
        assert area.contains(box)
        assert (box.left, box.top) == (area.left, area.top)
        assert not box.overlaps(header)
        assert not any(box.overlaps(f) for f in footers)
        assert pic._element.nvPicPr.cNvPr.get("descr") == entry.title
        assert "Exportiert" in slide.notes_slide.notes_text_frame.text


def test_unknown_layout_lists_available_names(samples, tmp_path):
    template, _ = samples
    with pytest.raises(ValueError, match="PBI Page"):
        build_deck(template, [], tmp_path / "deck.pptx", layout_name="Nope")


def test_preselects_largest_unobstructed_image_area(samples):
    template, _ = samples
    assert default_layout(open_template(template)).name == "PBI Page"


def test_graphic_inside_image_area_is_reported_and_avoided_by_preselection(samples):
    template, _ = samples
    prs, layout = _layout(template, "PBI Page")
    assert obstructions(layout, prs.slide_width, prs.slide_height) == []
    layout.shapes._spTree.append(_band(SLIDE_W // 2, HEADER_BOTTOM * 2, SLIDE_W // 4, HEADER_BOTTOM, "FF0000", 950))
    assert obstructions(layout, prs.slide_width, prs.slide_height) == ["Header band 950"]
    assert default_layout(prs).name != "PBI Page"
