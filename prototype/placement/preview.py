"""Draw a geometry preview of every slide in a Deck, without PowerPoint.

Draws master/layout decoration (header bands), title, footer and picture boxes at their real
positions, with the actual Page Image pasted in, and the Image Area outlined in red.
Usage: python preview.py DECK.pptx [OUT.png]
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from placement import Box, FOOTER_TYPES, TITLE_TYPES, image_area

WIDTH = 1000


def _scaled(box: Box, k: float):
    return [box.left * k, box.top * k, box.right * k, box.bottom * k]


def _decoration(shapes):
    for shape in shapes:
        if shape.is_placeholder or not shape.has_text_frame and shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        try:
            color = "#" + str(shape.fill.fore_color.rgb)
        except Exception:
            continue
        yield Box.of(shape), color


def render_slide(prs, slide) -> Image.Image:
    k = WIDTH / prs.slide_width
    img = Image.new("RGB", (WIDTH, int(prs.slide_height * k)), "white")
    draw = ImageDraw.Draw(img)
    layout = slide.slide_layout
    for box, color in [*_decoration(layout.slide_master.shapes), *_decoration(layout.shapes)]:
        draw.rectangle(_scaled(box, k), fill=color)
    for shape in slide.shapes:
        box = Box.of(shape)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            page = Image.open(io.BytesIO(shape.image.blob)).convert("RGB")
            w, h = int(box.width * k), int(box.height * k)
            img.paste(page.resize((w, h)), (int(box.left * k), int(box.top * k)))
        elif shape.is_placeholder and shape.placeholder_format.type in TITLE_TYPES:
            draw.rectangle(_scaled(box, k), outline="#1F3864", width=2)
            draw.text((box.left * k + 6, box.top * k + 6), shape.text_frame.text, fill="#1F3864")
        elif shape.is_placeholder and shape.placeholder_format.type in FOOTER_TYPES:
            draw.rectangle(_scaled(box, k), outline="#ED7D31", width=2)
            draw.text((box.left * k + 4, box.top * k + 4), shape.text_frame.text or "‹footer›", fill="#ED7D31")
        elif shape.has_text_frame:
            draw.text((box.left * k + 6, box.top * k + 6), shape.text_frame.text, fill="#C00000")
    area = image_area(layout, prs.slide_width, prs.slide_height)
    draw.rectangle(_scaled(area, k), outline="#FF0000", width=1)
    return img


def render_deck(deck: Path, out: Path) -> Path:
    prs = Presentation(deck)
    slides = [render_slide(prs, s) for s in prs.slides]
    w, h = slides[0].size
    cols = 2
    rows = (len(slides) + cols - 1) // cols
    sheet = Image.new("RGB", (cols * w + (cols + 1) * 12, rows * h + (rows + 1) * 12), "#404040")
    for i, im in enumerate(slides):
        sheet.paste(im, (12 + (i % cols) * (w + 12), 12 + (i // cols) * (h + 12)))
    sheet.save(out)
    return out


if __name__ == "__main__":
    deck = Path(sys.argv[1])
    print(render_deck(deck, Path(sys.argv[2]) if len(sys.argv) > 2 else deck.with_suffix(".preview.png")))
