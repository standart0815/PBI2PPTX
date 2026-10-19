"""Milestone 1: place Page Images into slides built from a Template.

Image Area rules (docs/DESIGN.md, "Slide composition"):
- If the Slide Layout has a picture placeholder, its box is the Image Area.
- Otherwise: full slide width, from the bottom of the title placeholder (the header)
  to the top of the highest footer placeholder, or the slide bottom.
- The Page Image is as large as possible inside the Image Area, anchored top-left,
  never cropped or stretched.
- Slides that come with the Template are dropped; the Deck has only page slides.
- Preselected Slide Layout: the largest Image Area that no decorative graphic reaches into.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence

from PIL import Image
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt

from potx import open_template

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
FOOTER_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}

LABELS = {
    "de": {
        "report": "Bericht",
        "page": "Seite",
        "link": "Link",
        "filters": "Filter",
        "no_filters": "keine",
        "exported": "Exportiert",
        "error": "Diese Seite konnte nicht gerendert werden.",
        "reason": "Grund",
    },
    "en": {
        "report": "Report",
        "page": "Page",
        "link": "Link",
        "filters": "Filters",
        "no_filters": "none",
        "exported": "Exported",
        "error": "This page could not be rendered.",
        "reason": "Reason",
    },
}


@dataclass(frozen=True)
class Box:
    left: int
    top: int
    width: int
    height: int

    @property
    def right(self) -> int:
        return self.left + self.width

    @property
    def bottom(self) -> int:
        return self.top + self.height

    def contains(self, other: "Box") -> bool:
        return (
            other.left >= self.left
            and other.top >= self.top
            and other.right <= self.right
            and other.bottom <= self.bottom
        )

    def overlaps(self, other: "Box") -> bool:
        return not (
            other.right <= self.left
            or other.left >= self.right
            or other.bottom <= self.top
            or other.top >= self.bottom
        )

    @classmethod
    def of(cls, shape) -> "Box":
        return cls(int(shape.left), int(shape.top), int(shape.width), int(shape.height))


@dataclass
class Entry:
    """One slide to produce. `image` is None for an Entry whose page failed to render."""

    title: str
    image: Optional[Path]
    report: str = ""
    link: str = ""
    filters: Sequence[str] = ()
    error: str = ""


def _ph_type(ph):
    return ph.placeholder_format.type


def image_area(layout, slide_width: int, slide_height: int) -> Box:
    pictures = [ph for ph in layout.placeholders if _ph_type(ph) == PP_PLACEHOLDER.PICTURE]
    if pictures:
        return Box.of(pictures[0])

    titles = [ph for ph in layout.placeholders if _ph_type(ph) in TITLE_TYPES]
    top = max((Box.of(t).bottom for t in titles), default=0)
    footer_tops = [
        int(ph.top) for ph in layout.placeholders if _ph_type(ph) in FOOTER_TYPES and int(ph.top) >= top
    ]
    bottom = min(footer_tops, default=int(slide_height))
    if bottom <= top:
        raise ValueError(f"Slide Layout '{layout.name}' leaves no room below the header")
    return Box(0, top, int(slide_width), bottom - top)


def fit_top_left(area: Box, image_width: int, image_height: int) -> Box:
    scale = min(area.width / image_width, area.height / image_height)
    return Box(area.left, area.top, int(image_width * scale), int(image_height * scale))


def layouts_by_name(prs) -> dict:
    return {layout.name: layout for layout in prs.slide_layouts}


def obstructions(layout, slide_width: int, slide_height: int) -> List[str]:
    """Names of decorative graphics (master or layout shapes) that reach into the Image Area."""
    area = image_area(layout, slide_width, slide_height)
    shapes = [*layout.slide_master.shapes, *layout.shapes]
    return [s.name for s in shapes if not s.is_placeholder and area.overlaps(Box.of(s))]


def default_layout(prs):
    """The layout with the largest Image Area that no decorative graphic reaches into."""
    best, best_size = None, 0
    for layout in prs.slide_layouts:
        if not any(_ph_type(ph) in TITLE_TYPES for ph in layout.placeholders):
            continue
        try:
            area = image_area(layout, prs.slide_width, prs.slide_height)
        except ValueError:
            continue
        if obstructions(layout, prs.slide_width, prs.slide_height):
            continue
        if area.width * area.height > best_size:
            best, best_size = layout, area.width * area.height
    if best is None:
        raise ValueError("Template has no layout with a title and an unobstructed Image Area")
    return best


def _drop_template_slides(prs) -> None:
    # Templates often ship example slides; a Deck contains only page slides.
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _copy_footers(layout, slide) -> None:
    # python-pptx does not copy date/footer/slide-number placeholders onto new slides.
    for ph in layout.placeholders:
        if _ph_type(ph) in FOOTER_TYPES:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))


def _remove_unused_placeholders(slide) -> None:
    # Body, content, table and picture placeholders cloned from the layout would show up
    # as empty "click to add" boxes; the slide keeps only its title.
    for ph in list(slide.placeholders):
        if _ph_type(ph) not in TITLE_TYPES:
            ph._element.getparent().remove(ph._element)


def _notes(entry: Entry, labels: dict, exported_at: datetime) -> str:
    filters = "; ".join(entry.filters) if entry.filters else labels["no_filters"]
    lines = [
        f"{labels['report']}: {entry.report}",
        f"{labels['page']}: {entry.title}",
        f"{labels['link']}: {entry.link}",
        f"{labels['filters']}: {filters}",
        f"{labels['exported']}: {exported_at.isoformat(timespec='minutes')}",
    ]
    if entry.error:
        lines.append(f"{labels['reason']}: {entry.error}")
    return "\n".join(lines)


def build_deck(
    template: Path,
    entries: List[Entry],
    out: Path,
    layout_name: Optional[str] = None,
    language: str = "de",
    exported_at: Optional[datetime] = None,
) -> Path:
    prs = open_template(template)
    if layout_name:
        layouts = layouts_by_name(prs)
        if layout_name not in layouts:
            raise ValueError(f"No Slide Layout '{layout_name}'. Available: {', '.join(layouts)}")
        layout = layouts[layout_name]
    else:
        layout = default_layout(prs)

    _drop_template_slides(prs)
    area = image_area(layout, prs.slide_width, prs.slide_height)
    labels = LABELS[language]
    exported_at = exported_at or datetime.now()

    for entry in entries:
        slide = prs.slides.add_slide(layout)
        _remove_unused_placeholders(slide)
        _copy_footers(layout, slide)
        if slide.shapes.title is not None:
            slide.shapes.title.text = entry.title

        if entry.image is not None:
            with Image.open(entry.image) as img:
                width, height = img.size
            box = fit_top_left(area, width, height)
            pic = slide.shapes.add_picture(str(entry.image), box.left, box.top, box.width, box.height)
            pic._element.nvPicPr.cNvPr.set("descr", entry.title)
        else:
            text = slide.shapes.add_textbox(area.left, area.top, area.width, area.height).text_frame
            text.word_wrap = True
            text.text = labels["error"]
            text.paragraphs[0].runs[0].font.size = Pt(24)
            reason = text.add_paragraph()
            reason.text = f"{labels['page']}: {entry.title} | {labels['reason']}: {entry.error}"
            reason.runs[0].font.size = Pt(14)

        slide.notes_slide.notes_text_frame.text = _notes(entry, labels, exported_at)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
