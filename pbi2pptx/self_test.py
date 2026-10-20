"""`PBI2PPTX --self-test`: checks that a packaged app has everything it needs to build a Deck.

Builds a small Deck (template, Page Image, speaker notes, Error Slide) entirely offline, so the build
scripts can run it right after packaging. It catches files the packager left out, such as
python-pptx's XML templates, which tests running from source never notice.
"""

from __future__ import annotations

import logging
import sys
import tempfile
from pathlib import Path

from PIL import Image
from pptx import Presentation


def run() -> int:
    from .deck import CapturedEntry, build_deck
    from .template import describe_template, save_as_potx

    try:
        with tempfile.TemporaryDirectory(prefix="pbi2pptx-self-test-") as folder:
            folder = Path(folder)
            template = folder / "template.potx"
            save_as_potx(Presentation(), template)  # python-pptx's own default template
            page = folder / "page.png"
            Image.new("RGB", (1280, 720), "white").save(page)
            info = describe_template(template.read_bytes(), template.name)
            entries = [
                CapturedEntry("Page", "https://app.powerbi.com/groups/me/reports/r/p", page, "Report", ["t/c = v"]),
                CapturedEntry("Missing", "https://app.powerbi.com/groups/me/reports/r/p", None, error="self-test"),
            ]
            deck = build_deck(template, entries, folder / "deck.pptx", layout_name=info.preselected)
            slides = list(Presentation(deck).slides)
            assert len(slides) == 2, f"expected 2 slides, got {len(slides)}"
            assert "t/c = v" in slides[0].notes_slide.notes_text_frame.text, "speaker notes missing"
        from . import static_dir

        assert (static_dir() / "index.html").exists(), f"interface missing in {static_dir()}"
    except Exception as e:  # noqa: BLE001 - reported, and the exit code fails the build
        message = f"SELF-TEST FAILED: {type(e).__name__}: {e}"
        logging.getLogger("pbi2pptx").error(message)  # the Windows app has no console: the log has it
        print(message, file=sys.stderr, flush=True)
        return 1
    message = "Self-test passed: template, Page Image, speaker notes, Error Slide and interface are all available."
    logging.getLogger("pbi2pptx").info(message)
    print(message, flush=True)
    return 0
