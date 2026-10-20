"""The Template: reading .potx files, Slide Layouts, the Image Area and Decorative Graphics.

Image Area rules (docs/DESIGN.md, "Slide composition"):
- If the Slide Layout has a picture placeholder, its box is the Image Area.
- Otherwise: full slide width, from the bottom of the title placeholder (the header)
  to the top of the highest footer placeholder, or the slide bottom.
- A Custom Image Area, set by the User for the whole Deck, replaces both.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, replace
from pathlib import Path
from typing import List, Optional

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
FOOTER_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}

_CT_PRESENTATION = "application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
_CT_TEMPLATE = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"


# --- .potx files (python-pptx only accepts the .pptx content type) ----------------------------


def _swap_main_content_type(data: bytes, old: str, new: str) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                content = content.replace(old.encode(), new.encode())
            dst.writestr(item, content)
    return out.getvalue()


class NotATemplate(ValueError):
    """The file can't be read as a PowerPoint template."""


def open_template(source: Path | bytes):
    """Open a .potx (or .pptx), from a path or from its bytes, as an editable Presentation."""
    data = source if isinstance(source, bytes) else Path(source).read_bytes()
    try:
        return Presentation(io.BytesIO(_swap_main_content_type(data, _CT_TEMPLATE, _CT_PRESENTATION)))
    except Exception as e:  # zipfile.BadZipFile, KeyError, lxml errors, python-pptx ValueError
        raise NotATemplate("This isn't a PowerPoint template (.potx or .pptx)") from e


def save_as_potx(prs, path: Path) -> None:
    buf = io.BytesIO()
    prs.save(buf)
    Path(path).write_bytes(_swap_main_content_type(buf.getvalue(), _CT_PRESENTATION, _CT_TEMPLATE))


# --- Geometry ---------------------------------------------------------------------------------


@dataclass(frozen=True)
class Box:
    left: int
    top: int
    width: int
    height: int

    @property
    def right(self) -> int:
        return self.left + self.width

    @property
    def bottom(self) -> int:
        return self.top + self.height

    @property
    def size(self) -> int:
        return self.width * self.height

    def contains(self, other: "Box") -> bool:
        return (
            other.left >= self.left
            and other.top >= self.top
            and other.right <= self.right
            and other.bottom <= self.bottom
        )

    def overlaps(self, other: "Box") -> bool:
        return not (
            other.right <= self.left
            or other.left >= self.right
            or other.bottom <= self.top
            or other.top >= self.bottom
        )

    @classmethod
    def of(cls, shape) -> "Box":
        return cls(int(shape.left), int(shape.top), int(shape.width), int(shape.height))


def fit_top_left(area: Box, image_width: int, image_height: int) -> Box:
    """As large as possible inside the area, anchored top-left, never cropped or stretched."""
    scale = min(area.width / image_width, area.height / image_height)
    return Box(area.left, area.top, int(image_width * scale), int(image_height * scale))


# --- Slide Layouts ----------------------------------------------------------------------------


def placeholder_type(ph):
    return ph.placeholder_format.type


def has_title(layout) -> bool:
    return any(placeholder_type(ph) in TITLE_TYPES for ph in layout.placeholders)


def image_area(layout, slide_width: int, slide_height: int) -> Box:
    pictures = [ph for ph in layout.placeholders if placeholder_type(ph) == PP_PLACEHOLDER.PICTURE]
    if pictures:
        return Box.of(pictures[0])

    titles = [ph for ph in layout.placeholders if placeholder_type(ph) in TITLE_TYPES]
    top = max((Box.of(t).bottom for t in titles), default=0)
    footer_tops = [
        int(ph.top) for ph in layout.placeholders if placeholder_type(ph) in FOOTER_TYPES and int(ph.top) >= top
    ]
    bottom = min(footer_tops, default=int(slide_height))
    if bottom <= top:
        raise ValueError(f"Slide Layout '{layout.name}' leaves no room below the header")
    return Box(0, top, int(slide_width), bottom - top)


EMU_PER_CM = 360_000
CUSTOM_AREA_SLACK = EMU_PER_CM // 20  # 0.5 mm: the interface rounds to hundredths of a cm


class OutsideSlide(ValueError):
    """A Custom Image Area that reaches outside the slide."""


def custom_image_area(box: Box, slide_width: int, slide_height: int) -> Box:
    """The Custom Image Area on this slide: refused if it reaches outside, trimmed to the slide edge
    if it only passes it by rounding."""
    if box.width <= 0 or box.height <= 0 or box.left < 0 or box.top < 0:
        raise OutsideSlide("The picture area needs a width and a height, and starts on the slide.")
    if box.right > slide_width + CUSTOM_AREA_SLACK or box.bottom > slide_height + CUSTOM_AREA_SLACK:
        raise OutsideSlide(
            f"The picture area reaches outside the slide ({slide_width / EMU_PER_CM:.2f} × {slide_height / EMU_PER_CM:.2f} cm)."
        )
    return Box(box.left, box.top, min(box.right, slide_width) - box.left, min(box.bottom, slide_height) - box.top)


def decorative_graphics_in_image_area(layout, slide_width: int, slide_height: int) -> List[str]:
    """Names of Decorative Graphics (non-placeholder shapes on the layout or master) inside the Image Area."""
    area = image_area(layout, slide_width, slide_height)
    shapes = [*layout.slide_master.shapes, *layout.shapes]
    return [s.name for s in shapes if not s.is_placeholder and area.overlaps(Box.of(s))]


@dataclass(frozen=True)
class LayoutInfo:
    """What the interface shows about one Slide Layout."""

    name: str
    image_area: Optional[Box]  # None: the layout can't hold a page slide
    graphics: tuple = ()  # (name, Box) of Decorative Graphics inside the Image Area; non-empty means a warning
    title: Optional[Box] = None
    footers: tuple = ()  # Box of each footer placeholder
    preselected: bool = False

    @property
    def usable(self) -> bool:
        return self.image_area is not None

    @property
    def decorative_graphics(self) -> tuple:
        return tuple(name for name, _ in self.graphics)


def _describe(layout, slide_width: int, slide_height: int) -> LayoutInfo:
    titles = [ph for ph in layout.placeholders if placeholder_type(ph) in TITLE_TYPES]
    footers = tuple(Box.of(ph) for ph in layout.placeholders if placeholder_type(ph) in FOOTER_TYPES)
    title = Box.of(titles[0]) if titles else None
    if title is None:
        return LayoutInfo(layout.name, None, footers=footers)
    try:
        area = image_area(layout, slide_width, slide_height)
    except ValueError:
        return LayoutInfo(layout.name, None, title=title, footers=footers)
    shapes = [*layout.slide_master.shapes, *layout.shapes]
    graphics = tuple((s.name, Box.of(s)) for s in shapes if not s.is_placeholder and area.overlaps(Box.of(s)))
    return LayoutInfo(layout.name, area, graphics, title, footers)


def describe_layouts(prs) -> List[LayoutInfo]:
    """Every Slide Layout, with its Image Area, warnings and which one is preselected."""
    infos = [_describe(layout, prs.slide_width, prs.slide_height) for layout in prs.slide_layouts]
    # Preselection: the largest Image Area that no Decorative Graphic reaches into; the first wins a tie.
    candidates = [i for i in infos if i.usable and not i.graphics]
    if not candidates:
        return infos
    best = max(candidates, key=lambda i: (i.image_area.size, -infos.index(i)))
    return [replace(i, preselected=i is best) for i in infos]


@dataclass(frozen=True)
class TemplateInfo:
    """A Template as the interface shows it."""

    file_name: str
    slide_width: int
    slide_height: int
    template_slides: int  # example slides in the Template; the Deck leaves them out
    layouts: tuple  # (LayoutInfo, …)

    @property
    def preselected(self) -> Optional[str]:
        return next((i.name for i in self.layouts if i.preselected), None)


def describe_template(source: Path | bytes, file_name: str) -> TemplateInfo:
    prs = open_template(source)
    return TemplateInfo(file_name, int(prs.slide_width), int(prs.slide_height), len(prs.slides), tuple(describe_layouts(prs)))


def preselected_layout(prs):
    for info, layout in zip(describe_layouts(prs), prs.slide_layouts):
        if info.preselected:
            return layout
    raise ValueError("Template has no layout with a title and an unobstructed Image Area")


def layout_by_name(prs, name: str):
    layouts = {layout.name: layout for layout in prs.slide_layouts}
    if name not in layouts:
        raise ValueError(f"No Slide Layout '{name}'. Available: {', '.join(layouts)}")
    return layouts[name]
