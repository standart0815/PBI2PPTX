"""Building the Deck: one slide per Entry, Page Images in the Image Area, notes and Error Slides.

Rules (docs/DESIGN.md, "Slide composition"):
- Slides that come with the Template are dropped; the Deck has only page slides.
- Of the layout's placeholders only the title and the footers stay on a slide.
- The Page Image is as large as possible inside the Image Area (or the Custom Image Area), anchored top-left.
- A page that could not be captured becomes an Error Slide; the Export carries on.
- With "Show Filters on slides", a Filter Box lists the Entry's Filters to the right of the page.
"""

from __future__ import annotations

import copy
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence

from PIL import Image
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR
from pptx.util import Cm, Pt

from .template import (
    FOOTER_TYPES,
    TITLE_TYPES,
    Box,
    custom_image_area,
    fit_top_left,
    image_area,
    layout_by_name,
    open_template,
    placeholder_type,
    preselected_layout,
)

LABELS = {
    "de": {
        "report": "Bericht",
        "page": "Seite",
        "link": "Link",
        "filters": "Filter",
        "no_filters": "keine",
        "exported": "Exportiert",
        "error": "Diese Seite konnte nicht erfasst werden.",
        "reason": "Grund",
        "filter_check": "Filter-Prüfung",
        "filter_box": "Filter",
    },
    "en": {
        "report": "Report",
        "page": "Page",
        "link": "Link",
        "filters": "Filters",
        "no_filters": "none",
        "exported": "Exported",
        "error": "This page could not be captured.",
        "reason": "Reason",
        "filter_check": "Filter check",
        "filter_box": "Filters",
    },
}


@dataclass
class CapturedEntry:
    """An Entry after capturing: its Page Image, or the reason there is none."""

    title: str
    page_url: str
    image: Optional[Path]
    report_name: str = ""
    filters: Sequence[str] = field(default_factory=tuple)
    error: str = ""
    filter_note: str = ""  # e.g. "not possible: the report hides its filter pane"
    filter_box: Sequence[str] = field(default_factory=tuple)  # lines for the Filter Box, e.g. "Country: Germany, France"


FILTER_BOX_GAP = Cm(0.3)
FILTER_BOX_SHARE = 0.20  # the Filter Box needs at least this share of the Image Area's width
FILTER_BOX_PT = 12


def filter_box_right_edge(layout, area: Box, slide_width: int, custom: bool = False) -> int:
    """Inside a Custom Image Area or a picture placeholder its edge; otherwise the title's left margin,
    mirrored on the right."""
    if custom or any(placeholder_type(ph) == PP_PLACEHOLDER.PICTURE for ph in layout.placeholders):
        return area.right
    titles = [ph for ph in layout.placeholders if placeholder_type(ph) in TITLE_TYPES]
    margin = int(titles[0].left) if titles and int(titles[0].left) < slide_width // 2 else 0
    return min(area.right, slide_width - margin)


def place_with_filter_box(area: Box, image_width: int, image_height: int, right_edge: int) -> tuple:
    """The Page Image and the Filter Box: the box goes where the page leaves room on the right, and
    the page is made just narrow enough when that room is less than the box needs."""
    need = FILTER_BOX_GAP + int(area.width * FILTER_BOX_SHARE)
    image = fit_top_left(area, image_width, image_height)
    if right_edge - image.right < need:
        narrowed = Box(area.left, area.top, right_edge - need - area.left, area.height)
        image = fit_top_left(narrowed, image_width, image_height)
    left = image.right + FILTER_BOX_GAP
    return image, Box(left, image.top, right_edge - left, area.bottom - image.top)


def _add_filter_box(slide, box: Box, heading: str, lines: Sequence[str]) -> None:
    shape = slide.shapes.add_textbox(box.left, box.top, box.width, box.height)
    shape.name = "Filter Box"
    text = shape.text_frame
    text.word_wrap = True
    text.vertical_anchor = MSO_ANCHOR.TOP
    text.text = heading
    text.paragraphs[0].runs[0].font.bold = True
    text.paragraphs[0].runs[0].font.size = Pt(FILTER_BOX_PT)
    for line in lines:  # the template's font and text colour, by not setting them
        paragraph = text.add_paragraph()
        paragraph.text = line
        paragraph.runs[0].font.size = Pt(FILTER_BOX_PT)


def _drop_template_slides(prs) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _remove_unused_placeholders(slide) -> None:
    # Body, content, table and picture placeholders cloned from the layout would show up
    # as empty "click to add" boxes.
    for ph in list(slide.placeholders):
        if placeholder_type(ph) not in TITLE_TYPES:
            ph._element.getparent().remove(ph._element)


def _copy_footers(layout, slide) -> None:
    # python-pptx does not copy date/footer/slide-number placeholders onto new slides.
    for ph in layout.placeholders:
        if placeholder_type(ph) in FOOTER_TYPES:
            slide.shapes._spTree.append(copy.deepcopy(ph._element))


def _notes(entry: CapturedEntry, labels: dict, exported_at: datetime) -> str:
    filters = "; ".join(entry.filters) if entry.filters else labels["no_filters"]
    lines = [
        f"{labels['report']}: {entry.report_name}",
        f"{labels['page']}: {entry.title}",
        f"{labels['link']}: {entry.page_url}",
        f"{labels['filters']}: {filters}",
        f"{labels['exported']}: {exported_at.isoformat(timespec='minutes')}",
    ]
    if entry.filter_note:
        lines.append(f"{labels['filter_check']}: {entry.filter_note}")
    if entry.error:
        lines.append(f"{labels['reason']}: {entry.error}")
    return "\n".join(lines)


def build_deck(
    template: Path | bytes,
    entries: List[CapturedEntry],
    out: Path,
    layout_name: Optional[str] = None,
    language: str = "de",
    exported_at: Optional[datetime] = None,
    custom_area: Optional[Box] = None,
) -> Path:
    prs = open_template(template)
    layout = layout_by_name(prs, layout_name) if layout_name else preselected_layout(prs)
    _drop_template_slides(prs)
    if custom_area is not None:
        area = custom_image_area(custom_area, int(prs.slide_width), int(prs.slide_height))
    else:
        area = image_area(layout, prs.slide_width, prs.slide_height)
    labels = LABELS[language]
    exported_at = exported_at or datetime.now()

    for entry in entries:
        slide = prs.slides.add_slide(layout)
        _remove_unused_placeholders(slide)
        _copy_footers(layout, slide)
        if slide.shapes.title is not None:
            slide.shapes.title.text = entry.title

        if entry.image is not None:
            with Image.open(entry.image) as img:
                width, height = img.size
            if entry.filter_box:
                right_edge = filter_box_right_edge(layout, area, int(prs.slide_width), custom_area is not None)
                box, filter_box = place_with_filter_box(area, width, height, right_edge)
            else:
                box, filter_box = fit_top_left(area, width, height), None
            pic = slide.shapes.add_picture(str(entry.image), box.left, box.top, box.width, box.height)
            pic._element.nvPicPr.cNvPr.set("descr", entry.title)
            if filter_box is not None:
                _add_filter_box(slide, filter_box, labels["filter_box"], entry.filter_box)
        else:
            text = slide.shapes.add_textbox(area.left, area.top, area.width, area.height).text_frame
            text.word_wrap = True
            text.text = labels["error"]
            text.paragraphs[0].runs[0].font.size = Pt(24)
            reason = text.add_paragraph()
            reason.text = f"{labels['page']}: {entry.title} | {labels['reason']}: {entry.error}"
            reason.runs[0].font.size = Pt(14)

        slide.notes_slide.notes_text_frame.text = _notes(entry, labels, exported_at)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
