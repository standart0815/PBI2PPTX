from datetime import datetime

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pbi2pptx.deck import CapturedEntry, build_deck
from pbi2pptx.template import FOOTER_TYPES, TITLE_TYPES, Box, image_area
from tests.fixtures.samples import HEADER_BOTTOM, SLIDE_W

URL = "https://app.powerbi.com/groups/me/reports/0e0e0e0e-0000-0000-0000-000000000004/cacea11bd78861e6a8a2"


def _entries(pages):
    return [CapturedEntry(p.stem, URL, p, "Example", ["financials/Country = Germany"]) for p in pages] + [
        CapturedEntry("broken", URL, None, "Example", error="No report page at this address")
    ]


@pytest.mark.parametrize("layout_name", ["PBI Page", "PBI Page (Picture)"])
def test_page_slides_never_cover_header_or_footers(sample_template, sample_pages, tmp_path, layout_name):
    entries = _entries(sample_pages)
    prs = Presentation(build_deck(sample_template, entries, tmp_path / "deck.pptx", layout_name=layout_name))
    slides = list(prs.slides)
    assert len(slides) == len(entries)  # the template's cover slide is dropped

    header = Box(0, 0, SLIDE_W, HEADER_BOTTOM)
    for slide, entry in zip(slides, entries):
        assert slide.shapes.title.text == entry.title
        assert {ph.placeholder_format.type for ph in slide.placeholders} <= TITLE_TYPES | FOOTER_TYPES
        footers = [Box.of(ph) for ph in slide.placeholders if ph.placeholder_format.type in FOOTER_TYPES]
        assert len(footers) == 3
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        if entry.image is None:
            assert not pictures
            continue
        (pic,) = pictures
        box = Box.of(pic)
        area = image_area(slide.slide_layout, prs.slide_width, prs.slide_height)
        assert area.contains(box)
        assert (box.left, box.top) == (area.left, area.top)
        assert not box.overlaps(header)
        assert not any(box.overlaps(f) for f in footers)
        assert pic._element.nvPicPr.cNvPr.get("descr") == entry.title


def test_uses_preselected_layout_when_none_is_chosen(sample_template, sample_pages, tmp_path):
    prs = Presentation(build_deck(sample_template, _entries(sample_pages[:1]), tmp_path / "deck.pptx"))
    assert prs.slides[0].slide_layout.name == "PBI Page"


def test_error_slide_names_page_and_reason(sample_template, tmp_path):
    entry = CapturedEntry("Umsatz", URL, None, "Example", error="No report page at this address")
    slide = Presentation(build_deck(sample_template, [entry], tmp_path / "deck.pptx")).slides[0]
    text = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
    assert "Diese Seite konnte nicht erfasst werden." in text
    assert "Umsatz" in text and "No report page at this address" in text


@pytest.mark.parametrize("language, label", [("de", "Exportiert: 2026-10-19T18:43"), ("en", "Exported: 2026-10-19T18:43")])
def test_notes_trace_each_slide_back_in_the_render_language(sample_template, sample_pages, tmp_path, language, label):
    entry = CapturedEntry("Example", URL, sample_pages[0], "Example", ["financials/Country = Germany"])
    deck = build_deck(sample_template, [entry], tmp_path / "deck.pptx", language=language, exported_at=datetime(2026, 10, 19, 18, 43))
    notes = Presentation(deck).slides[0].notes_slide.notes_text_frame.text
    assert label in notes
    assert URL in notes and "financials/Country = Germany" in notes


# --- Filter Box -------------------------------------------------------------------------------

from pptx.util import Cm  # noqa: E402

from pbi2pptx.deck import FILTER_BOX_GAP, FILTER_BOX_SHARE, place_with_filter_box  # noqa: E402
from tests.fixtures.samples import FOOTER_TOP, MARGIN  # noqa: E402

LINES = ("Country: Germany, France", "Segment: Government")


def _boxes(slide):
    picture = next(Box.of(s) for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    filter_box = next((s for s in slide.shapes if s.name == "Filter Box"), None)
    return picture, filter_box


def test_filter_box_uses_the_room_the_page_leaves():
    area = Box(0, 0, int(Cm(33.9)), int(Cm(13.5)))  # "Tabelle" in Firmenvorlage
    right = int(Cm(33.9 - 2.5))
    image, box = place_with_filter_box(area, 2560, 1440, right)
    assert image == Box(0, 0, int(Cm(13.5) * 2560 / 1440), int(Cm(13.5)))  # the page keeps its full size
    assert box.left == image.right + FILTER_BOX_GAP and box.right == right
    assert box.width >= area.width * FILTER_BOX_SHARE


def test_page_is_narrowed_when_the_room_is_too_small():
    area = Box(0, 0, int(Cm(33.9)), int(Cm(13.5)))
    right = int(Cm(31.4))
    image, box = place_with_filter_box(area, 3840, 1080, right)  # a wide page would fill the width
    assert box.right == right
    assert box.width == pytest.approx(area.width * FILTER_BOX_SHARE, abs=2)
    assert image.right + FILTER_BOX_GAP == box.left
    assert image.width / image.height == pytest.approx(3840 / 1080, rel=1e-3)


@pytest.mark.parametrize("layout_name", ["PBI Page", "PBI Page (Picture)"])
def test_filter_box_on_slides(sample_template, sample_pages, tmp_path, layout_name):
    entries = [
        CapturedEntry("with Filters", URL, sample_pages[0], "Example", filter_box=LINES),
        CapturedEntry("without Filters", URL, sample_pages[0], "Example"),
        CapturedEntry("failed", URL, None, "Example", error="x", filter_box=LINES),
    ]
    prs = Presentation(build_deck(sample_template, entries, tmp_path / "deck.pptx", layout_name=layout_name, language="en"))
    with_filters, without_filters, failed = prs.slides
    area = image_area(with_filters.slide_layout, prs.slide_width, prs.slide_height)

    picture, filter_box = _boxes(with_filters)
    box = Box.of(filter_box)
    assert filter_box.text_frame.text.splitlines() == ["Filters", *LINES]
    assert area.contains(picture) and area.contains(box)
    assert box.left >= picture.right and box.top == picture.top
    assert not box.overlaps(Box(0, 0, prs.slide_width, HEADER_BOTTOM))
    assert box.bottom <= FOOTER_TOP
    if layout_name == "PBI Page":
        assert box.right == prs.slide_width - MARGIN  # the title's margin, mirrored
    else:
        assert box.right == area.right  # inside the picture placeholder

    unfiltered_picture, none = _boxes(without_filters)
    assert none is None
    assert unfiltered_picture.width >= picture.width  # only slides with Filters make room
    assert all(s.name != "Filter Box" for s in failed.shapes)


# --- Custom Image Area ------------------------------------------------------------------------

from PIL import Image  # noqa: E402

from pbi2pptx.template import OutsideSlide, fit_top_left  # noqa: E402

CUSTOM = Box(int(Cm(2)), int(Cm(4)), int(Cm(20)), int(Cm(10)))


def _size(path):
    with Image.open(path) as img:
        return img.size


def test_custom_image_area_holds_page_filter_box_and_error_text(sample_template, sample_pages, tmp_path):
    entries = [
        CapturedEntry("with Filters", URL, sample_pages[0], "Example", filter_box=LINES),
        CapturedEntry("without Filters", URL, sample_pages[0], "Example"),
        CapturedEntry("failed", URL, None, "Example", error="x"),
    ]
    prs = Presentation(build_deck(sample_template, entries, tmp_path / "deck.pptx", language="en", custom_area=CUSTOM))
    with_filters, without_filters, failed = prs.slides

    picture, filter_box = _boxes(with_filters)
    assert (picture.left, picture.top) == (CUSTOM.left, CUSTOM.top)
    assert CUSTOM.contains(picture) and CUSTOM.contains(Box.of(filter_box))
    assert Box.of(filter_box).right == CUSTOM.right

    alone, _ = _boxes(without_filters)
    assert alone == fit_top_left(CUSTOM, *_size(sample_pages[0]))

    (text,) = [Box.of(s) for s in failed.shapes if s.has_text_frame and "could not be captured" in s.text_frame.text]
    assert text == CUSTOM


def test_custom_image_area_outside_the_slide_is_refused(sample_template, sample_pages, tmp_path):
    entry = CapturedEntry("Example", URL, sample_pages[0], "Example")
    beyond = Box(int(Cm(20)), 0, int(Cm(20)), int(Cm(5)))  # the slide is 33.9 cm wide
    with pytest.raises(OutsideSlide, match="outside the slide"):
        build_deck(sample_template, [entry], tmp_path / "deck.pptx", custom_area=beyond)
