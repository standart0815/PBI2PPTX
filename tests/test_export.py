"""The Export end to end: window service + fake Power BI page + sample template → Deck."""

import asyncio
import io

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from pbi2pptx.export import EntryState, Exporter, ExportRefused, ExportState
from pbi2pptx.page_url import Filter
from pbi2pptx.selection import Entry, Selection
from tests.test_power_bi_window import MISSING, REPORT, WS_PAGE
from tests.test_window_service import service_for

APP_PAGE = f"https://app.powerbi.com/groups/me/apps/a1a1a1a1-0000-0000-0000-000000000005/reports/{REPORT}/aaaa0000000000000002"
GERMANY = Filter.parse("financials/Country=Germany")
TYPO = Filter.parse("financials/Land=Germany")


def entry(url, title, *filters):
    return Entry(url, title, "Beispielbericht", title, list(filters))


async def run_export(tmp_path, selection, template, before=None, today=None, hidden=False):
    service, _ = service_for(tmp_path)
    await service.start()
    exporter = Exporter(service, today=lambda: today) if today else Exporter(service)
    try:
        if before:
            await before(service)
        exporter.start(selection, template, hidden=hidden)
        while exporter.running:
            await asyncio.sleep(0.2)
        return exporter, service.language
    finally:
        await service.stop()


def slides(deck: bytes):
    return list(Presentation(io.BytesIO(deck)).slides)


def text_of(slide) -> str:
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def test_exports_every_entry_and_turns_failures_into_error_slides(tmp_path, sample_template):
    selection = Selection(
        "Monthly pack: EMEA",
        [
            entry(WS_PAGE, "Umsatz"),
            entry(WS_PAGE, "Umsatz Deutschland", GERMANY),
            entry(WS_PAGE, "Umsatz mit Tippfehler", TYPO),
            entry(MISSING, "Gelöschte Seite"),
            entry(APP_PAGE, "Regionen (App)", GERMANY),
        ],
    )
    exporter, _ = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes()))

    progress = exporter.progress
    assert progress.state is ExportState.DONE
    assert [e.state for e in progress.entries] == [
        EntryState.CAPTURED,
        EntryState.CAPTURED,
        EntryState.ERROR_SLIDE,
        EntryState.ERROR_SLIDE,
        EntryState.CAPTURED,
    ]
    assert "Power BI did not apply: financials/Land = Germany" in progress.entries[2].message

    file_name, deck = exporter.deck()
    assert file_name.startswith("Monthly pack EMEA ") and file_name.endswith(".pptx")
    deck_slides = slides(deck)
    assert [s.shapes.title.text for s in deck_slides] == [e.slide_title for e in selection.entries]
    pictures = [any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in s.shapes) for s in deck_slides]
    assert pictures == [True, True, False, False, True]
    assert "Power BI hat diese Filter nicht angewendet: financials/Land = Germany" in text_of(deck_slides[2])
    assert "Keine Berichtsseite unter dieser Adresse" in text_of(deck_slides[3])


def test_hidden_filter_pane_is_exported_but_marked_not_verified(tmp_path, sample_template):
    async def hide_pane(service):
        await service._window.page.evaluate("() => localStorage.setItem('fake-no-pane', '1')")

    selection = Selection("Pack", [entry(WS_PAGE, "Umsatz Deutschland", GERMANY)])
    exporter, _ = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes(), before=hide_pane))
    (item,) = exporter.progress.entries
    assert item.state is EntryState.NOT_VERIFIED
    (slide,) = slides(exporter.deck()[1])
    assert "Filter-Prüfung: nicht möglich" in slide.notes_slide.notes_text_frame.text


def test_render_language_reopens_the_window_in_that_language(tmp_path, sample_template):
    selection = Selection("Pack", [entry(WS_PAGE, "Sales", TYPO)], render_language="en-US")
    exporter, language = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes()))
    assert language == "en-US"
    (slide,) = slides(exporter.deck()[1])
    assert "Power BI did not apply these Filters" in text_of(slide)  # Error Slide in English


def test_cancel_stops_without_a_deck(tmp_path, sample_template):
    async def go():
        service, _ = service_for(tmp_path)
        await service.start()
        exporter = Exporter(service)
        try:
            exporter.start(Selection("Pack", [entry(WS_PAGE, "A"), entry(WS_PAGE, "B"), entry(WS_PAGE, "C")]), sample_template.read_bytes())
            exporter.cancel()
            while exporter.running:
                await asyncio.sleep(0.2)
            return exporter
        finally:
            await service.stop()

    exporter = asyncio.run(go())
    assert exporter.progress.state is ExportState.CANCELLED
    assert exporter.deck() is None


def test_refuses_to_start_without_what_it_needs(tmp_path, sample_template):
    async def go():
        service, _ = service_for(tmp_path)
        exporter = Exporter(service)
        with pytest.raises(ExportRefused, match="template"):
            exporter.start(Selection("Pack", [entry(WS_PAGE, "A")]), None)
        with pytest.raises(ExportRefused, match="no Entries"):
            exporter.start(Selection("Pack", []), sample_template.read_bytes())
        # The window isn't open: the Export starts, then fails with a message for the User.
        exporter.start(Selection("Pack", [entry(WS_PAGE, "A")]), sample_template.read_bytes())
        with pytest.raises(ExportRefused, match="already running"):
            exporter.start(Selection("Pack", [entry(WS_PAGE, "A")]), sample_template.read_bytes())
        while exporter.running:
            await asyncio.sleep(0.1)
        return exporter.progress

    progress = asyncio.run(go())
    assert progress.state is ExportState.FAILED and "isn't open" in progress.error


def test_show_filters_puts_a_filter_box_on_filtered_slides(tmp_path, sample_template):
    selection = Selection("Pack", [entry(WS_PAGE, "DE", GERMANY), entry(WS_PAGE, "All")], show_filters=True)
    exporter, _ = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes()))
    with_filters, without = slides(exporter.deck()[1])
    (filter_box,) = [s for s in with_filters.shapes if s.name == "Filter Box"]
    assert filter_box.text_frame.text.splitlines() == ["Filter", "Country: Germany"]
    assert all(s.name != "Filter Box" for s in without.shapes)



# --- Date Filters -------------------------------------------------------------------------------

from datetime import date  # noqa: E402

from pbi2pptx.dates import DateRule  # noqa: E402
from pbi2pptx.page_url import DateFilter  # noqa: E402

LAST_YEAR = DateFilter("Calendar", "Date", DateRule.START_OF_LAST_YEAR, DateRule.END_OF_LAST_YEAR)
SINCE_JUNE = DateFilter("Calendar", "Date", start=date(2013, 6, 1))


def notes_of(slide) -> str:
    return slide.notes_slide.notes_text_frame.text


@pytest.mark.parametrize(
    "language, box, note",
    [
        ("de-DE", ["Filter", "Date: 01.01.2013 – 31.12.2013", "Date: ab 01.06.2013"],
         "Calendar/Date: 01.01.2013 (Anfang letzten Jahres) – 31.12.2013 (Ende letzten Jahres)"),
        ("en-US", ["Filters", "Date: 01/01/2013 – 12/31/2013", "Date: from 06/01/2013"],
         "Calendar/Date: 01/01/2013 (Start of last year) – 12/31/2013 (End of last year)"),
    ],
)
def test_date_rules_are_worked_out_when_the_export_starts(tmp_path, sample_template, language, box, note):
    selection = Selection("Pack", [entry(WS_PAGE, "Last year", LAST_YEAR, SINCE_JUNE)], render_language=language, show_filters=True)
    exporter, _ = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes(), today=date(2014, 1, 15)))
    (item,) = exporter.progress.entries
    assert item.state is EntryState.CAPTURED  # Power BI showed a card for each end
    (slide,) = slides(exporter.deck()[1])
    (filter_box,) = [s for s in slide.shapes if s.name == "Filter Box"]
    assert filter_box.text_frame.text.splitlines() == box
    assert note in notes_of(slide)


def test_impossible_pages_become_error_slides(tmp_path, sample_template):
    this_month_so_far = DateFilter("Calendar", "Date", DateRule.START_OF_THIS_MONTH, DateRule.YESTERDAY)
    many = [Filter("financials", "Segment", (f"S{i}",)) for i in range(9)]
    long = Filter("financials", "Country", tuple(f"Country number {i}" for i in range(80)))
    selection = Selection(
        "Pack",
        [
            entry(WS_PAGE, "Empty on the 1st", this_month_so_far),
            entry(WS_PAGE, "Eleven conditions", LAST_YEAR, *many),
            entry(WS_PAGE, "Too long", long),
        ],
    )
    exporter, _ = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes(), today=date(2026, 11, 1)))
    assert [e.state for e in exporter.progress.entries] == [EntryState.ERROR_SLIDE] * 3
    empty, eleven, too_long = (text_of(s) for s in slides(exporter.deck()[1]))
    assert "Calendar/Date" in empty and "01.11.2026" in empty and "31.10.2026" in empty
    assert "11" in eleven and "10" in eleven
    assert "2.000" in too_long or "2000" in too_long


def test_a_hidden_export_gives_the_same_deck(tmp_path, sample_template):
    selection = Selection("Hidden", [entry(WS_PAGE, "Umsatz"), entry(WS_PAGE, "Umsatz Deutschland", GERMANY)])
    exporter, _ = asyncio.run(run_export(tmp_path, selection, sample_template.read_bytes(), hidden=True))
    assert exporter.progress.state is ExportState.DONE
    assert [e.state for e in exporter.progress.entries] == [EntryState.CAPTURED, EntryState.CAPTURED]
    pictures = [s for slide in slides(exporter.deck()[1]) for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pictures) == 2

