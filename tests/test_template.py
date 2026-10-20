import pytest
from pptx.enum.shapes import PP_PLACEHOLDER

from pbi2pptx.template import (
    Box,
    decorative_graphics_in_image_area,
    describe_layouts,
    fit_top_left,
    image_area,
    layout_by_name,
    open_template,
    preselected_layout,
)
from tests.fixtures.samples import FOOTER_TOP, HEADER_BOTTOM, MARGIN, SLIDE_W, TITLE_TOP, band


def _layout(template, name):
    prs = open_template(template)
    return prs, layout_by_name(prs, name)


def test_opens_potx_files(sample_template):
    prs = open_template(sample_template)
    assert len(prs.slides) == 1


def test_default_image_area_is_full_width_between_header_and_footers(sample_template):
    prs, layout = _layout(sample_template, "PBI Page")
    assert image_area(layout, prs.slide_width, prs.slide_height) == Box(0, HEADER_BOTTOM, SLIDE_W, FOOTER_TOP - HEADER_BOTTOM)


def test_picture_placeholder_overrides_image_area(sample_template):
    prs, layout = _layout(sample_template, "PBI Page (Picture)")
    pic = next(ph for ph in layout.placeholders if ph.placeholder_format.type == PP_PLACEHOLDER.PICTURE)
    assert image_area(layout, prs.slide_width, prs.slide_height) == Box.of(pic)


def test_fit_is_height_limited_and_anchored_top_left():
    area = Box(0, 1000, 12_000, 5_000)
    box = fit_top_left(area, 1600, 900)
    assert (box.left, box.top, box.height) == (0, 1000, 5_000)
    assert box.width < area.width


def test_fit_uses_full_width_when_height_allows():
    box = fit_top_left(Box(0, 1000, 12_000, 5_000), 3200, 900)
    assert (box.left, box.top, box.width) == (0, 1000, 12_000)


def test_fit_never_distorts():
    box = fit_top_left(Box(0, 0, 12_000, 5_000), 1920, 1440)
    assert box.width / box.height == pytest.approx(1920 / 1440, rel=1e-3)


def test_preselects_largest_unobstructed_image_area(sample_template):
    assert preselected_layout(open_template(sample_template)).name == "PBI Page"


def test_describes_every_layout_for_the_interface(sample_template):
    infos = {i.name: i for i in describe_layouts(open_template(sample_template))}
    assert infos["PBI Page"].preselected and infos["PBI Page"].usable
    assert not infos["PBI Page (Picture)"].preselected
    assert [i.name for i in infos.values() if i.preselected] == ["PBI Page"]


def test_decorative_graphic_inside_image_area_is_reported_and_avoided(sample_template):
    prs, layout = _layout(sample_template, "PBI Page")
    assert decorative_graphics_in_image_area(layout, prs.slide_width, prs.slide_height) == []
    layout.shapes._spTree.append(band(SLIDE_W // 2, HEADER_BOTTOM * 2, SLIDE_W // 4, HEADER_BOTTOM, "FF0000", 950))
    assert decorative_graphics_in_image_area(layout, prs.slide_width, prs.slide_height) == ["Header band 950"]
    assert preselected_layout(prs).name != "PBI Page"


def test_unknown_layout_lists_available_names(sample_template):
    with pytest.raises(ValueError, match="PBI Page"):
        layout_by_name(open_template(sample_template), "Nope")


def test_describes_a_template_from_its_bytes(sample_template):
    from pbi2pptx.template import describe_template

    info = describe_template(sample_template.read_bytes(), "Sample.potx")
    assert (info.file_name, info.template_slides, info.preselected) == ("Sample.potx", 1, "PBI Page")
    page = next(i for i in info.layouts if i.name == "PBI Page")
    assert page.title == Box(int(MARGIN), int(TITLE_TOP), int(SLIDE_W - 2 * MARGIN), int(HEADER_BOTTOM - TITLE_TOP))
    assert len(page.footers) == 3 and page.graphics == ()


def test_graphics_in_the_image_area_carry_their_boxes(sample_template):
    from pbi2pptx.template import describe_layouts

    prs, layout = _layout(sample_template, "PBI Page")
    layout.shapes._spTree.append(band(SLIDE_W // 2, HEADER_BOTTOM * 2, SLIDE_W // 4, HEADER_BOTTOM, "FF0000", 951))
    info = next(i for i in describe_layouts(prs) if i.name == "PBI Page")
    ((name, box),) = info.graphics
    assert name == "Header band 951" and box.left == SLIDE_W // 2
    assert info.decorative_graphics == ("Header band 951",)


def test_refuses_files_that_are_not_templates(tmp_path):
    from pbi2pptx.template import NotATemplate

    with pytest.raises(NotATemplate):
        open_template(b"PK\x03\x04 not really a zip")
    with pytest.raises(NotATemplate):
        open_template(b"%PDF-1.7")


def test_custom_image_area_must_stay_on_the_slide():
    from pbi2pptx.template import CUSTOM_AREA_SLACK, OutsideSlide, custom_image_area

    slide_w, slide_h = 1000_000, 600_000
    assert custom_image_area(Box(10, 20, 500, 300), slide_w, slide_h) == Box(10, 20, 500, 300)
    # passing the edge by rounding only is trimmed back to the edge
    assert custom_image_area(Box(0, 0, slide_w + CUSTOM_AREA_SLACK, slide_h), slide_w, slide_h) == Box(0, 0, slide_w, slide_h)
    for outside in (Box(0, 0, slide_w + CUSTOM_AREA_SLACK + 1, 10), Box(-1, 0, 10, 10), Box(0, 0, 0, 10)):
        with pytest.raises(OutsideSlide):
            custom_image_area(outside, slide_w, slide_h)
